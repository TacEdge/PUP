#!/usr/bin/env python3
"""
Build the one-page ELDA Pathway Overview as a landscape branded .docx.
Four courses side by side (Lead Teams, Lead Leaders, Lead Systems, Command)
with scope, duration, learners, aim, developmental logic, learning outcomes
and next step, drawn from the revised course data sheets.
"""

import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from PIL import Image

import build_elda_data_sheet as cds

# ----------------------------------------------------------------- CONFIG ---
LOGO_FILE          = "./assets/nz-army-logo.png"
OUTPUT_DOCX        = "./output/elda-pathway-overview.docx"
PROTECTIVE_MARKING = "UNCLASSIFIED"
DOCUMENT_REFERENCE = "CDS A18011, A18008, A18010, A18009"
FOOTER_REFERENCE   = "ACS 2026"
DATE               = "September 2026"
ORIGINATOR         = "NZ Army Leadership Centre | Army Command School"

TITLE         = "Experiential Leadership Development Activity Pathway"
SUBTITLE_LINE = "Experiential Leadership Development Activities: Course Overview"
STATUS        = "Draft for Validation"
FOOTER_LEFT   = "NZALC | ELDA Pathway Overview"

ARMY_RED      = "C62026"
DARKEST_HOUR  = "000000"
RUAPEHU_WHITE = "FFFFFF"
SWAMP_GREEN   = "002516"
WAIOURU_HILLS = "A89662"
MOAWHANGO     = "CDD2B7"
PALE_GREEN    = "EEF1E5"
GRID_GREY     = "D9D9D2"

FONT = "Arial"

# ---------------------------------------------------------------- CONTENT ---

COURSES = [
    dict(
        name="ELDA Lead Teams",
        code="A18011",
        duration="6 Training Days",
        unit="Individual / Team",
        scope="Leading a team",
        learners="Regular Force personnel selected for the A1530 Promotion "
                 "All Corps RF JNCO Course.",
        aim="Develop individual leadership effectiveness through small-team "
            "leadership under pressure.",
        logic=["LEAD SELF", "LEAD TEAM", "REFLECT", "DEVELOP"],
        outcomes=[
            "Apply strategies to maintain effective performance under pressure",
            "Apply effective leadership behaviours within a small team "
            "under challenging conditions",
            "Evaluate personal leadership performance and develop an "
            "individual leadership development plan",
        ],
        next_step="ELDA Lead Leaders.",
    ),
    dict(
        name="ELDA Lead Leaders",
        code="A18008",
        duration="6 Training Days",
        unit="Other Leaders",
        scope="Leading among leaders",
        learners="Regular Force Officers (2LT and LT); Regular Force NCOs "
                 "accepted onto the A1531 SNCO Promotion Course.",
        aim="Develop leadership effectiveness through self-awareness, "
            "reflection and deliberate adaptation of leadership behaviour.",
        logic=["UNDERSTAND SELF", "UNDERSTAND OTHERS", "ADAPT BEHAVIOUR",
               "DEVELOP DELIBERATELY"],
        outcomes=[
            "Demonstrate Army ethos and values through leadership "
            "behaviour under challenging conditions",
            "Evaluate personal leadership behaviour and its impact on others",
            "Adapt leadership behaviour to improve effectiveness",
            "Develop an individual leadership development strategy",
        ],
        next_step="ELDA Lead Systems or ELDA Command, as appropriate to "
                  "career pathway.",
    ),
    dict(
        name="ELDA Lead Systems",
        code="A18010",
        duration="7 Training Days",
        unit="System",
        scope="Leading across a system",
        learners="SNCOs and WOs accepted onto the A1532 All Corps Warrant "
                 "Officer Course; Captains preparing for promotion to Major.",
        aim="Develop leadership effectiveness within complex systems through "
            "understanding interdependencies, exercising influence and "
            "adapting leadership behaviour.",
        logic=["UNDERSTAND THE SYSTEM", "NAVIGATE INTERDEPENDENCE",
               "INFLUENCE OUTCOMES", "ADAPT & DEVELOP"],
        outcomes=[
            "Demonstrate Army ethos, values and sound judgement within "
            "complex environments",
            "Analyse relationships and interdependencies within a complex "
            "system",
            "Influence others to achieve outcomes beyond direct authority",
            "Evaluate and adapt leadership behaviour to improve system "
            "effectiveness",
        ],
        next_step="ELDA Command, as appropriate to career pathway.",
    ),
    dict(
        name="ELDA Command",
        code="A18009",
        duration="8–10 Training Days",
        unit="Command Collective / Organisation",
        scope="Leading collectively at command level",
        learners="Established or forming command teams and leadership groups, "
                 "including unit HQs, command teams, principal staff and "
                 "subordinate command teams.",
        aim="Develop collective command effectiveness through shared "
            "understanding, cohesion, collective judgement and deliberate "
            "creation of the conditions for organisational performance.",
        logic=["UNDERSTAND EACH OTHER", "BUILD COHESION",
               "EXERCISE COLLECTIVE JUDGEMENT", "CREATE THE CONDITIONS"],
        outcomes=[
            "Develop shared understanding within the command team",
            "Strengthen cohesion and effectiveness across the command team",
            "Exercise collective judgement in complex and demanding "
            "conditions",
            "Create the conditions for organisational performance",
        ],
        next_step="Continued command-team development and application within "
                  "command and senior leadership appointments.",
    ),
]

SCOPE_CAPTION = "Increasing scope, interdependence and organisational consequence"

CLOSING = (
    "The ELDA pathway deliberately expands the scope and complexity of "
    "leadership: from leading self and a small team, to understanding one's "
    "impact on other leaders, to influencing outcomes across an "
    "interconnected system, and ultimately to developing the collective "
    "effectiveness of command teams that create the conditions for "
    "organisational performance."
)

# ------------------------------------------------------------ XML helpers ---


def _pPr(p):
    return p._p.get_or_add_pPr()


def set_shading(p, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    _pPr(p).append(shd)


def set_border(p, side, color, sz, space=4):
    pPr = _pPr(p)
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pBdr.append(el)


def force_font(style_or_run, name):
    style_or_run.font.name = name
    rPr = style_or_run.font.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for attr in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
        rFonts.attrib.pop(qn(f"w:{attr}"), None)
    rFonts.set(qn("w:cs"), name)


def force_color(style_or_run, hexval):
    style_or_run.font.color.rgb = RGBColor.from_string(hexval)
    rPr = style_or_run.font.element.get_or_add_rPr()
    color = rPr.find(qn("w:color"))
    if color is not None:
        for attr in ("themeColor", "themeShade", "themeTint"):
            color.attrib.pop(qn(f"w:{attr}"), None)


def strip_style_rpr(style):
    rPr = style.element.find(qn("w:rPr"))
    if rPr is not None:
        style.element.remove(rPr)


def add_field(p, code, placeholder=""):
    runs = []
    for kind in ("begin", None, "separate", "text", "end"):
        r = p.add_run(placeholder if kind == "text" else "")
        if kind in ("begin", "separate", "end"):
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            r._r.append(fld)
        elif kind is None:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = f" {code} "
            r._r.append(instr)
        runs.append(r)
    return runs


def run(p, text, size, bold=False, color=DARKEST_HOUR, italic=False):
    r = p.add_run(text)
    force_font(r, FONT)
    r.font.size = size
    r.bold = bold
    r.italic = italic
    force_color(r, color)
    return r


def cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_margins(cell, top=40, bottom=40, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("bottom", bottom),
                      ("left", left), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def table_borders(table, color, sz=4, outer=None):
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), outer if (outer and not side.startswith("inside")) else color)
        borders.append(el)
    table._tbl.tblPr.append(borders)


# ------------------------------------------------------------- doc set-up ---

TEXT_WIDTH_CM = 29.7 - 3.6


def set_landscape(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Cm(1.1)
    section.bottom_margin = Cm(1.1)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)


def apply_styles(doc):
    normal = doc.styles["Normal"]
    strip_style_rpr(normal)
    force_font(normal, FONT)
    normal.font.size = Pt(9)
    force_color(normal, DARKEST_HOUR)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.04


# ------------------------------------------------------- headers & footers --


def marking_paragraph(container):
    p = container.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, PROTECTIVE_MARKING, Pt(10), bold=True)
    return p


def add_header_footer(section, footer_left=FOOTER_LEFT,
                      text_width_cm=TEXT_WIDTH_CM):
    marking_paragraph(section.header)
    marking_paragraph(section.footer)
    info = section.footer.add_paragraph()
    info.paragraph_format.space_before = Pt(2)
    info.paragraph_format.tab_stops.add_tab_stop(
        Cm(text_width_cm / 2), WD_TAB_ALIGNMENT.CENTER)
    info.paragraph_format.tab_stops.add_tab_stop(
        Cm(text_width_cm), WD_TAB_ALIGNMENT.RIGHT)
    run(info, footer_left, Pt(8.5))
    run(info, "\t" + FOOTER_REFERENCE, Pt(8.5))
    run(info, "\tPage ", Pt(8.5))
    for r in add_field(info, "PAGE", "1"):
        force_font(r, FONT)
        r.font.size = Pt(8.5)
    run(info, " of ", Pt(8.5))
    for r in add_field(info, "NUMPAGES", "1"):
        force_font(r, FONT)
        r.font.size = Pt(8.5)


# -------------------------------------------------------------- letterhead --


def add_letterhead(doc, with_logo=True, title_style=None):
    """Compact layered opening for the landscape page, then the pathway
    band: the four courses in sequence, which is the page's organising
    idea."""
    if with_logo:
        logo = Image.open(LOGO_FILE).convert("RGBA")
        logo = logo.crop(logo.getchannel("A").getbbox())
        buf = io.BytesIO()
        logo.save(buf, "PNG")
        buf.seek(0)
        doc.add_picture(buf, width=Mm(24))
        doc.paragraphs[-1].paragraph_format.space_after = Pt(5)

    k = doc.add_paragraph()
    k.paragraph_format.space_after = Pt(0)
    r = run(k, "NEW ZEALAND ARMY LEADERSHIP CENTRE", Pt(8.5), bold=True, color=SWAMP_GREEN)
    cds.letterspace(r, 40)

    title_p = doc.add_paragraph(style=title_style) if title_style \
        else doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(3)
    set_border(title_p, "bottom", ARMY_RED, 18, space=4)
    run(title_p, TITLE, Pt(17), bold=True)

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_before = Pt(3)
    meta_p.paragraph_format.space_after = Pt(0)
    run(meta_p, cds.ORIGINATOR_LONG, Pt(9), color=SWAMP_GREEN)
    run(meta_p, "      ", Pt(9))
    for i, (label, value) in enumerate([("Reference", DOCUMENT_REFERENCE),
                                        ("Date", DATE), ("Status", STATUS)]):
        if i:
            run(meta_p, "   ·   ", Pt(8.5), color=WAIOURU_HILLS)
        run(meta_p, f"{label}: ", Pt(8.5), bold=True, color=SWAMP_GREEN)
        run(meta_p, value, Pt(8.5), bold=(label == "Status"),
            color=ARMY_RED if label == "Status" else DARKEST_HOUR)

    # The pathway band: course sequence and scope, the first thing to read.
    band = doc.add_paragraph()
    band.paragraph_format.space_before = Pt(5)
    band.paragraph_format.space_after = Pt(0)
    for i, course in enumerate(COURSES):
        if i:
            run(band, "   →   ", Pt(12), bold=True, color=ARMY_RED)
        r = run(band, course["name"].replace("ELDA ", "").upper(), Pt(12),
                bold=True, color=SWAMP_GREEN)
        cds.letterspace(r, 20)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(2)
    r = run(cap, SCOPE_CAPTION.upper() + "  →", Pt(7), bold=True,
            color=WAIOURU_HILLS)
    cds.letterspace(r, 25)
    return title_p


# ------------------------------------------------------------------ table ---

LABEL_CM = 2.9
ROWS = ["course", "unit", "scope", "learners", "aim", "logic", "outcomes", "next"]
LABELS = {"unit": "", "scope": "Scope", "learners": "Learners", "aim": "Aim",
          "logic": "Developmental logic", "outcomes": "Learning outcomes",
          "next": "Next step"}


def render_pathway_table(doc, text_width_cm=TEXT_WIDTH_CM):
    col_cm = (text_width_cm - LABEL_CM) / len(COURSES)
    table = doc.add_table(rows=len(ROWS), cols=len(COURSES) + 1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_borders(table, GRID_GREY, outer=WAIOURU_HILLS)
    for c_idx, col in enumerate(table.columns):
        col.width = Cm(LABEL_CM if c_idx == 0 else col_cm)

    for r_idx, key in enumerate(ROWS):
        row = table.rows[r_idx]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        for c_idx in range(len(COURSES) + 1):
            cell = row.cells[c_idx]
            cell.width = Cm(LABEL_CM if c_idx == 0 else col_cm)
            pad = 12 if key == "unit" else (45 if key in ("scope", "logic") else 35)
            cell_margins(cell, top=pad, bottom=pad)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.04

            if c_idx == 0:
                if key == "unit":
                    cell_shading(cell, WAIOURU_HILLS)
                    r = run(p, "SCOPE", Pt(7), bold=True, color=RUAPEHU_WHITE)
                    cds.letterspace(r, 30)
                    continue
                cell_shading(cell, SWAMP_GREEN if key == "course" else MOAWHANGO)
                if key != "course":
                    run(p, LABELS[key], Pt(8.5), bold=True, color=SWAMP_GREEN)
                continue

            course = COURSES[c_idx - 1]
            if key == "course":
                cell_shading(cell, SWAMP_GREEN)
                run(p, course["name"], Pt(11), bold=True, color=RUAPEHU_WHITE)
                p.add_run().add_break()
                run(p, f'{course["code"]}  ·  {course["duration"]}', Pt(8),
                    color=MOAWHANGO)
            elif key == "unit":
                # Thin scope bar: what is changing across the pathway.
                cell_shading(cell, WAIOURU_HILLS)
                label = course["unit"].upper()
                if c_idx < len(COURSES):
                    label += "   →"
                r = run(p, label, Pt(7.5), bold=True, color=RUAPEHU_WHITE)
                cds.letterspace(r, 25)
            elif key == "scope":
                cell_shading(cell, MOAWHANGO)
                set_border(p, "left", ARMY_RED, 18, space=4)
                run(p, course["scope"], Pt(10.5), bold=True, color=SWAMP_GREEN)
            elif key == "logic":
                # Same four-stage architecture in every column; read down for
                # the course's journey, across for the widening scope.
                cell_shading(cell, PALE_GREEN)
                p.paragraph_format.line_spacing = 1.0
                for i, stage in enumerate(course["logic"]):
                    if i:
                        p.add_run().add_break()
                        run(p, "▼", Pt(7), color=ARMY_RED)
                        p.add_run().add_break()
                    r = run(p, stage, Pt(8.5), bold=True, color=SWAMP_GREEN)
                    cds.letterspace(r, 8)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif key == "outcomes":
                for i, lo in enumerate(course["outcomes"]):
                    if i:
                        p.add_run().add_break()
                    run(p, f"LO 1.{i + 1}  ", Pt(7.5), bold=True, color=SWAMP_GREEN)
                    run(p, lo, Pt(7.5))
            else:
                run(p, course["next_step" if key == "next" else key], Pt(8))
    return table


def render_closing(doc):
    """Key message panel: dark green, white text, red rule."""
    close_p = doc.add_paragraph()
    close_p.paragraph_format.space_before = Pt(6)
    close_p.paragraph_format.space_after = Pt(0)
    close_p.paragraph_format.left_indent = Cm(0.3)
    close_p.paragraph_format.right_indent = Cm(0.3)
    set_shading(close_p, SWAMP_GREEN)
    set_border(close_p, "left", ARMY_RED, 28, space=10)
    r = run(close_p, "KEY MESSAGE", Pt(7.5), bold=True, color=MOAWHANGO)
    cds.letterspace(r, 30)
    close_p.add_run().add_break()
    run(close_p, CLOSING, Pt(8.5), bold=True, color=RUAPEHU_WHITE)
    return close_p


# ----------------------------------------------------------------- finish ---


def finish(doc, used_style_ids=("Normal", "DefaultParagraphFont",
                                "TableNormal", "TableGrid", "NoList",
                                "Header", "Footer")):
    styles_el = doc.styles.element
    for st in list(styles_el.findall(qn("w:style"))):
        if st.get(qn("w:styleId")) not in used_style_ids:
            styles_el.remove(st)
    for rFonts in styles_el.iter(qn("w:rFonts")):
        for attr in ("ascii", "hAnsi", "cs"):
            rFonts.set(qn(f"w:{attr}"), FONT)
        for attr in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme",
                     "eastAsia"):
            rFonts.attrib.pop(qn(f"w:{attr}"), None)
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)


def build():
    doc = Document()
    set_landscape(doc.sections[0])
    apply_styles(doc)
    add_header_footer(doc.sections[0])
    add_letterhead(doc)
    render_pathway_table(doc)
    render_closing(doc)
    finish(doc)
    doc.core_properties.title = TITLE
    doc.core_properties.author = ORIGINATOR
    doc.save(OUTPUT_DOCX)
    print(f"Saved {OUTPUT_DOCX}")


if __name__ == "__main__":
    build()
