#!/usr/bin/env python3
"""
Build the NZALC ELDA Course Portfolio: one branded .docx that compiles

  1. a cover page,
  2. a contents page (Word TOC field, resolved by convert_to_pdf.py),
  3. the one-page landscape ELDA Pathway Overview, and
  4. the four revised Course Data Sheets in pathway order
     (Lead Teams, Lead Leaders, Lead Systems, Command).

Content and styling come from build_elda_data_sheet.py and
build_elda_pathway.py, so the portfolio always matches the standalone
documents.  Run after any change to a data-sheet source:

    python3 build_elda_portfolio.py
    python3 convert_to_pdf.py output/elda-course-portfolio.docx \\
                              output/elda-course-portfolio.pdf
"""

from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.shared import Cm, Mm, Pt

import build_elda_data_sheet as cds
import build_elda_pathway as pathway

OUTPUT_DOCX = "./output/elda-course-portfolio.docx"
TITLE = "NZALC ELDA Course Portfolio"
SUBTITLE = ("Experiential Leadership Development Activities: "
            "Pathway Overview and Course Data Sheets")
REFERENCE = "Course Data Sheets A18011, A18008, A18010, A18009"
FOOTER_LEFT = "NZALC ELDA Course Portfolio"

# Pathway order, not course-code order.
SHEET_ORDER = ["lead-teams", "lead-leaders", "lead-systems", "command"]


def set_portrait(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)


def build():
    doc = cds.new_document()
    cds.SECTION_HEADING_STYLE = "Heading 2"   # data-sheet sections nest under each sheet
    cds.add_header_footer(doc.sections[0], FOOTER_LEFT)

    # ---- 1. cover -------------------------------------------------------
    cds.add_logo(width_mm=48, space_after=60)
    t = doc.add_paragraph()
    t.paragraph_format.space_after = Pt(4)
    cds.add_text_runs(t, TITLE, base_font=cds.FONT_HEAD, size=Pt(28), bold=True)
    s = doc.add_paragraph()
    s.paragraph_format.space_after = Pt(6)
    cds.set_border(s, "bottom", cds.ARMY_RED, 24, space=10)
    cds.add_text_runs(s, SUBTITLE, base_font=cds.FONT_HEAD, size=Pt(13),
                      bold=True, color=cds.SWAMP_GREEN)
    o = doc.add_paragraph()
    o.paragraph_format.space_before = Pt(10)
    cds.add_text_runs(o, cds.ORIGINATOR, base_font=cds.FONT_HEAD, size=Pt(11),
                      color=cds.SWAMP_GREEN)
    m = doc.add_paragraph()
    for i, (label, value) in enumerate([("Reference", REFERENCE), ("Date", cds.DATE)]):
        if i:
            cds.add_text_runs(m, "      ", base_font=cds.FONT_HEAD)
        cds.add_text_runs(m, f"{label}  ", base_font=cds.FONT_HEAD, size=Pt(9.5),
                          bold=True, color=cds.SWAMP_GREEN)
        cds.add_text_runs(m, value, base_font=cds.FONT_HEAD, size=Pt(9.5))

    intro = doc.add_paragraph()
    intro.paragraph_format.space_before = Pt(36)
    cds.add_text_runs(intro, (
        "This portfolio brings together the revised course data sheets for "
        "the four Experiential Leadership Development Activity (ELDA) "
        "courses delivered by the New Zealand Army Leadership Centre, "
        "together with a one-page overview of the pathway they form."))
    for line in [
        "ELDA Pathway Overview",
        "A18011 ELDA Lead Teams",
        "A18008 ELDA Lead Leaders",
        "A18010 ELDA Lead Systems",
        "A18009 ELDA Command",
    ]:
        cds.add_bullet(line)

    # ---- 2. contents ----------------------------------------------------
    toc_head = doc.add_paragraph()
    toc_head.paragraph_format.page_break_before = True
    toc_head.paragraph_format.space_after = Pt(12)
    cds.add_text_runs(toc_head, "Contents", base_font=cds.FONT_HEAD,
                      size=Pt(16), bold=True)
    toc_p = doc.add_paragraph()
    cds.add_field(toc_p, r'TOC \o "1-2" \h \z \u',
                  "Right-click and choose Update Field to refresh the contents.")

    # ---- 3. pathway overview (landscape section) ------------------------
    land = doc.add_section(WD_SECTION.NEW_PAGE)
    pathway.set_landscape(land)
    # Landscape footer tabs need the wider text width.
    land.footer.is_linked_to_previous = False
    land.header.is_linked_to_previous = False
    cds.add_header_footer(land, FOOTER_LEFT, text_width_cm=pathway.TEXT_WIDTH_CM)

    pathway.add_letterhead(doc, with_logo=False, title_style="Heading 1")
    pathway.render_pathway_table(doc)
    pathway.render_closing(doc)

    # ---- 4. course data sheets (portrait section) -----------------------
    port = doc.add_section(WD_SECTION.NEW_PAGE)
    set_portrait(port)
    port.footer.is_linked_to_previous = False
    port.header.is_linked_to_previous = False
    cds.add_header_footer(port, FOOTER_LEFT)

    for n, key in enumerate(SHEET_ORDER):
        sheet = cds.SHEETS[key]
        title = cds.add_letterhead(sheet["title"], cds.SUBTITLE_LINE,
                                   sheet["reference"], with_logo=False,
                                   title_style="Heading 1")
        if n:
            title.paragraph_format.page_break_before = True
        with open(sheet["source"], encoding="utf-8") as fh:
            cds.render_markdown(fh.read().splitlines())

    cds.finish(TITLE)
    doc.save(OUTPUT_DOCX)
    print(f"Saved {OUTPUT_DOCX}")


if __name__ == "__main__":
    build()
