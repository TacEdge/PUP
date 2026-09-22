#!/usr/bin/env python3
"""
Finish the ATG storyboard deck: add the blank template slide from the
ATG_STORYBOARD layout (so every field shows PowerPoint's click-to-add prompt),
move it to the front, and draw the photo number markers on it.

    node build_atg_storyboard.js && python3 build_atg_storyboard_finish.py
"""

import copy

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Inches, Pt

OUT = "output/atg-storyboard-template.pptx"

# must match build_atg_storyboard.js
GRID = dict(x=4.15, y=1.3, w=5.5, h=3.25, gap=0.1)
PH_W = (GRID["w"] - GRID["gap"]) / 2
PH_H = (GRID["h"] - GRID["gap"]) / 2
PHOTOS = [(GRID["x"] + (i % 2) * (PH_W + GRID["gap"]), GRID["y"] + (i // 2) * (PH_H + GRID["gap"])) for i in range(4)]


def marker(slide, x, y, n):
    d = 0.26
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.09), Inches(y + 0.09), Inches(d), Inches(d))
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor(0, 0, 0)
    sp.line.fill.background()
    sp.shadow.inherit = False
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.name = "Arial"
    r.font.color.rgb = RGBColor(255, 255, 255)
    sp.name = f"Photo {n} number"


def main():
    prs = Presentation(OUT)
    layout = next(l for l in prs.slide_layouts if l.name == "ATG_STORYBOARD")
    # pptxgenjs writes picture slots as untyped content placeholders and drops
    # the shrink-to-fit flag: type the photo slots and restore autofit here.
    for ph in layout.placeholders:
        elm = ph._element.nvSpPr.nvPr.ph
        if elm.get("type") is None:
            elm.set("type", "pic")
        elif elm.get("type") in ("body", "title"):
            ph.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for s in prs.slides:
        for ph in s.placeholders:
            if ph.has_text_frame:
                ph.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    slide = prs.slides.add_slide(layout)          # clones every placeholder, empty
    for i, (x, y) in enumerate(PHOTOS, 1):
        marker(slide, x, y, i)
    # move the new slide to the front
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    sldIdLst.remove(ids[-1])
    sldIdLst.insert(0, ids[-1])
    prs.save(OUT)
    names = [s.name for s in slide.placeholders]
    print(f"template slide added with {len(names)} placeholders: {', '.join(names)}")


if __name__ == "__main__":
    main()
