#!/usr/bin/env python3
"""
Combat Mindset Conditioning: product description as an editable Word
document in the shared house style, set in the black and grey palette.

    python3 build_cmc_product_description_docx.py
    /usr/bin/python3 convert_to_pdf.py output/combat-mindset-conditioning.docx output/combat-mindset-conditioning.pdf
"""

from docx.shared import Pt

import build_elda_data_sheet as cds

SOURCE = "./combat-mindset-conditioning.md"
OUTPUT_DOCX = "./output/combat-mindset-conditioning.docx"
TITLE = "Combat Mindset Conditioning"
TAG = "An Army training product within the Army Combat Mindset System"
FOOTER_LEFT = "Combat Mindset Conditioning | Draft for discussion"


# black and grey palette for this document: the pipeline's greens are
# swapped before any style is built (the constants are read at call time)
CHARCOAL = "222222"
LIGHT_GREY = "E6E6E2"
PALE_GREY = "F3F3F0"
MID_GREY = "8A8A8A"


def build():
    cds.SWAMP_GREEN = CHARCOAL
    cds.MOAWHANGO = LIGHT_GREY
    cds.PALE_GREEN = PALE_GREY
    cds.WAIOURU_HILLS = MID_GREY
    cds.DATE = "23 September 2026"
    cds.ORIGINATOR_LONG = "Army Command School"
    doc = cds.new_document()
    doc.styles["Normal"].font.size = Pt(9.5)
    cds.add_header_footer(doc.sections[0], FOOTER_LEFT)
    cds.add_title_block(None, TITLE, TAG, reference=None, status="Draft for discussion", date_only=True)
    with open(SOURCE, encoding="utf-8") as fh:
        cds.render_markdown(fh.read().splitlines())
    for para in doc.paragraphs:
        if para.style is not None and para.style.name == cds.SECTION_HEADING_STYLE:
            para.paragraph_format.space_before = Pt(8)
            para.paragraph_format.space_after = Pt(3)
        elif para.text.strip() and para.paragraph_format.space_after is None:
            para.paragraph_format.space_after = Pt(3)
            para.paragraph_format.line_spacing = 1.08
        if not para.text.strip() and para.runs == []:
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.line_spacing = Pt(4)
    # a lead-in paragraph stays with the table it introduces
    from docx.text.paragraph import Paragraph
    body = doc.element.body
    children = list(body.iterchildren())
    for elm, nxt in zip(children, children[1:]):
        if elm.tag.endswith("}p") and nxt.tag.endswith("}tbl"):
            Paragraph(elm, doc).paragraph_format.keep_with_next = True
    for table in doc.tables:
        # keep short tables on one page: rows hold on to the next; long
        # tables may break rather than leave a page mostly empty
        for row in (table.rows[:-1] if len(table.rows) <= 6 else []):
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.keep_with_next = True
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        if run.font.size == Pt(9):
                            run.font.size = Pt(8.5)
                    para.paragraph_format.line_spacing = 1.05
    cds.finish(TITLE)
    doc.save(OUTPUT_DOCX)
    print(f"Saved {OUTPUT_DOCX}")


if __name__ == "__main__":
    build()
