#!/usr/bin/env python3
"""
Build a revised ELDA Course Data Sheet as a compact branded .docx.
Pick the sheet by name:  python3 build_elda_data_sheet.py lead-leaders
(default: lead-teams). Letterhead-style first page (no cover, no TOC); body typeset
from SOURCE_FILE on the established brand system.

Source mini-syntax (a superset of the SOP builder's):
  ## / ### / ####   section, sub-section and learning-outcome headings
  | a | b |          two-column field tables (separator row ignored)
  >! text           review-note callout
  >> LABEL | text   one box in the learning-architecture chain
  * text / 1. text  bullets and numbered points (four-space '    * ' nests)
"""

import io
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from PIL import Image

# ----------------------------------------------------------------- CONFIG ---
SHEETS = {
    "lead-teams": dict(
        source="./elda-lead-teams-data-sheet.md",
        output="./output/elda-lead-teams-data-sheet.docx",
        code="A18011", name="ELDA Lead Teams",
        title="A18011 ELDA Lead Teams",
        reference="SOLO Course Data Sheet A18011, AL 3.0",
        footer_left="NZALC | ELDA Lead Teams Course Data Sheet",
    ),
    "lead-systems": dict(
        source="./elda-lead-systems-data-sheet.md",
        output="./output/elda-lead-systems-data-sheet.docx",
        code="A18010", name="ELDA Lead Systems",
        title="A18010 ELDA Lead Systems",
        reference="SOLO Course Data Sheet A18010, AL 3.0",
        footer_left="NZALC | ELDA Lead Systems Course Data Sheet",
    ),
    "command": dict(
        source="./elda-command-data-sheet.md",
        output="./output/elda-command-data-sheet.docx",
        code="A18009", name="ELDA Command",
        title="A18009 ELDA Command",
        reference="SOLO Course Data Sheet A18009, AL 2.5",
        footer_left="NZALC | ELDA Command Course Data Sheet",
    ),
    "lead-leaders": dict(
        source="./elda-lead-leaders-data-sheet.md",
        output="./output/elda-lead-leaders-data-sheet.docx",
        code="A18008", name="ELDA Lead Leaders",
        title="A18008 ELDA Lead Leaders",
        reference="SOLO Course Data Sheet A18008, AL 3.1",
        footer_left="NZALC | ELDA Lead Leaders Course Data Sheet",
    ),
}
LOGO_FILE          = "./assets/nz-army-logo.png"
PROTECTIVE_MARKING = "UNCLASSIFIED"
FOOTER_REFERENCE   = "ACS 2026"
DATE               = "September 2026"
ORIGINATOR         = "NZ Army Leadership Centre | Army Command School"
ORIGINATOR_LONG    = "New Zealand Army Leadership Centre | Army Command School"
SUBTITLE_LINE      = "Proposed Revised Course Data Sheet"
STATUS             = "Draft for Validation"

# Style used for "## " section headings; a compiled portfolio demotes it.
SECTION_HEADING_STYLE = "Heading 1"
doc = None   # bound by new_document() or by an importing builder

ARMY_RED      = "C62026"
DARKEST_HOUR  = "000000"
RUAPEHU_WHITE = "FFFFFF"
SWAMP_GREEN   = "002516"
WAIOURU_HILLS = "A89662"
MOAWHANGO     = "CDD2B7"
PALE_GREEN    = "EEF1E5"   # learning-outcome blocks, logic row
GRID_GREY     = "D9D9D2"   # internal table rules
MID_GREY      = "8A8A8A"   # de-emphasised locator steps

# The developmental pathway, in order; the locator line under each course
# title marks the current course.
PATHWAY = ["LEAD TEAMS", "LEAD LEADERS", "LEAD SYSTEMS", "COMMAND"]
PATHWAY_INDEX = {"ELDA Lead Teams": 0, "ELDA Lead Leaders": 1,
                 "ELDA Lead Systems": 2, "ELDA Command": 3}

FONT_HEAD = "Arial"
FONT_BODY = "Arial"

# ------------------------------------------------------------ XML helpers ---


def _pPr(p):
    return p._p.get_or_add_pPr()


def set_shading(p, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    _pPr(p).append(shd)


def set_border(p, side, color, sz, space=4):
    pPr = _pPr(p)
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pBdr.append(el)


def force_font(style_or_run, name):
    style_or_run.font.name = name
    rPr = style_or_run.font.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for attr in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
        rFonts.attrib.pop(qn(f"w:{attr}"), None)
    rFonts.set(qn("w:cs"), name)


def force_color(style_or_run, hexval):
    style_or_run.font.color.rgb = RGBColor.from_string(hexval)
    rPr = style_or_run.font.element.get_or_add_rPr()
    color = rPr.find(qn("w:color"))
    if color is not None:
        color.attrib.pop(qn("w:themeColor"), None)
        color.attrib.pop(qn("w:themeShade"), None)
        color.attrib.pop(qn("w:themeTint"), None)


def strip_style_rpr(style):
    rPr = style.element.find(qn("w:rPr"))
    if rPr is not None:
        style.element.remove(rPr)


def add_field(p, code, placeholder=None):
    r1 = p.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    r1._r.append(fld)
    r2 = p.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {code} "
    r2._r.append(instr)
    r3 = p.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3._r.append(sep)
    r4 = p.add_run(placeholder if placeholder else "")
    r5 = p.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r5._r.append(end)
    return [r1, r2, r3, r4, r5]


def mark_update_fields(doc):
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)


BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def add_text_runs(p, text, base_font=FONT_BODY, size=None, bold=False,
                  italic=False, color=None):
    idx = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > idx:
            _run(p, text[idx:m.start()], base_font, size, bold, italic, color)
        _run(p, m.group(1), base_font, size, True, italic, color)
        idx = m.end()
    if idx < len(text):
        _run(p, text[idx:], base_font, size, bold, italic, color)


def _run(p, text, font, size, bold, italic, color):
    run = p.add_run(text)
    force_font(run, font)
    if size:
        run.font.size = size
    run.bold = bold
    run.italic = italic
    if color:
        force_color(run, color)
    return run


def letterspace(run, twentieths=30):
    """Track out a run slightly (value in twentieths of a point)."""
    rPr = run._r.get_or_add_rPr()
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:val"), str(twentieths))
    rPr.append(sp)


# Rendering state set by render_markdown: body paragraphs under a learning
# outcome are indented to sit within its block, and Notes are subordinated.
_context = {"lo": False, "notes": False}


# ------------------------------------------------------------- doc set-up ---

TEXT_WIDTH_CM = 16.4


def new_document():
    """A4 portrait document with the brand Normal / Heading styles."""
    global doc
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    apply_styles(doc)
    return doc


def apply_styles(document):
    normal = document.styles["Normal"]
    strip_style_rpr(normal)
    force_font(normal, FONT_BODY)
    normal.font.size = Pt(10.5)
    force_color(normal, DARKEST_HOUR)
    nf = normal.paragraph_format
    nf.line_spacing = 1.12
    nf.space_after = Pt(5)
    nf.space_before = Pt(0)
    nf.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for name in ("Heading 1", "Heading 2"):
        h = document.styles[name]
        strip_style_rpr(h)
        force_font(h, FONT_HEAD)
        h.font.bold = True
        h.font.size = Pt(12)
        force_color(h, SWAMP_GREEN)
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        h.paragraph_format.keep_with_next = True


def add_section_heading(text):
    """Major section: green heading, 18 pt above, 8 pt below (one rule
    everywhere)."""
    _context["lo"] = False
    _context["notes"] = False
    p = doc.add_paragraph(style=SECTION_HEADING_STYLE)
    add_text_runs(p, text, base_font=FONT_HEAD, size=Pt(12.5), bold=True,
                  color=SWAMP_GREEN)
    return p


def add_body(text):
    p = doc.add_paragraph()
    if _context["lo"]:
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.space_after = Pt(4)
    add_text_runs(p, text)
    return p


def add_sub_heading(text, level):
    """### sub-section headings and #### learning-outcome blocks."""
    _context["lo"] = False
    _context["notes"] = False
    if level == 4:
        return add_outcome_block(text)

    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    add_text_runs(p, text, base_font=FONT_HEAD, size=Pt(10.5), bold=True,
                  color=SWAMP_GREEN)
    if text.strip() == "Notes":
        _context["notes"] = True
    return p


def add_outcome_block(text):
    """A learning outcome: pale green block with a green left rule, the LO
    number as a small label and the title beneath it."""
    label, _, title = text.partition(":")
    if not title:
        label, title = "", text
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(0.55)
    p.paragraph_format.right_indent = Cm(0.3)
    set_shading(p, PALE_GREEN)
    set_border(p, "left", SWAMP_GREEN, 14, space=10)
    if label:
        r = _run(p, label.strip().upper(), FONT_HEAD, Pt(8.5), True, False,
                 SWAMP_GREEN)
        letterspace(r, 20)
        p.add_run().add_break()
    add_text_runs(p, title.strip(), base_font=FONT_HEAD, size=Pt(11),
                  bold=True, color=DARKEST_HOUR)
    _context["lo"] = True
    return p


def _cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("bottom", bottom),
                      ("left", left), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def _table_borders(table, color, sz=4, outer=None):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), outer if (outer and not side.startswith("inside")) else color)
        borders.append(el)
    tblPr.append(borders)


def _cell_valign_top(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), "top")
    tcPr.append(va)


FIELD_COL_CM = 4.6


def add_wide_table(rows):
    """General table (three or more columns): dark header row, shaded first
    column, equal widths for the remaining columns."""
    ncols = max(len(r) for r in rows)
    rows = [tuple(r) + ("",) * (ncols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.autofit = False
    _table_borders(table, GRID_GREY, outer=WAIOURU_HILLS)
    first = Cm(3.6)
    rest = Cm((TEXT_WIDTH_CM - 3.6) / (ncols - 1))
    widths = [first] + [rest] * (ncols - 1)
    for col, width in zip(table.columns, widths):
        col.width = width
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.width = widths[c_idx]
            _cell_margins(cell, top=60, bottom=60, left=100, right=100)
            _cell_valign_top(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            if r_idx == 0:
                _cell_shading(cell, SWAMP_GREEN)
                add_text_runs(p, text, base_font=FONT_HEAD, size=Pt(9),
                              bold=True, color=RUAPEHU_WHITE)
            elif c_idx == 0:
                _cell_shading(cell, MOAWHANGO)
                add_text_runs(p, text, base_font=FONT_HEAD, size=Pt(9),
                              bold=True, color=SWAMP_GREEN)
            else:
                for k, part in enumerate(text.split("<br>")):
                    if k:
                        p.add_run().add_break()
                    add_text_runs(p, part.strip(), size=Pt(9))
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)
    return table


def add_field_table(rows):
    """Two-column table: shaded field labels on the left, entries on the
    right.  Light internal rules, top-aligned cells, generous padding."""
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=len(body) + 1, cols=2)
    table.autofit = False
    _table_borders(table, GRID_GREY, outer=WAIOURU_HILLS)
    widths = (Cm(FIELD_COL_CM), Cm(TEXT_WIDTH_CM - FIELD_COL_CM))
    for col, width in zip(table.columns, widths):
        col.width = width
    for r_idx, (label, value) in enumerate([header] + body):
        is_header = r_idx == 0
        for c_idx, text in enumerate((label, value)):
            cell = table.cell(r_idx, c_idx)
            cell.width = widths[c_idx]
            _cell_margins(cell, top=60, bottom=60, left=110, right=110)
            _cell_valign_top(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            if is_header:
                _cell_shading(cell, SWAMP_GREEN)
                add_text_runs(p, text, base_font=FONT_HEAD, size=Pt(9),
                              bold=True, color=RUAPEHU_WHITE)
            elif c_idx == 0:
                _cell_shading(cell, MOAWHANGO)
                add_text_runs(p, text, base_font=FONT_HEAD, size=Pt(9),
                              bold=True, color=SWAMP_GREEN)
            else:
                for k, part in enumerate(text.split("<br>")):
                    if k:
                        p.add_run().add_break()
                    add_text_runs(p, part.strip(), size=Pt(9.5))
    for r_idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        # Short tables stay in one piece: every row but the last keeps with
        # the next, and the preceding heading already keeps with the table.
        if len(table.rows) <= 8 and r_idx < len(table.rows) - 1:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)
    return table


def add_chain(boxes):
    """Learning-architecture chain: shaded boxes joined by down arrows."""
    for i, (label, text) in enumerate(boxes):
        if i:
            arrow = doc.add_paragraph()
            arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
            arrow.paragraph_format.space_before = Pt(0)
            arrow.paragraph_format.space_after = Pt(0)
            a = arrow.add_run("▼")
            force_font(a, FONT_HEAD)
            a.font.size = Pt(9)
            force_color(a, ARMY_RED)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_shading(p, MOAWHANGO)
        set_border(p, "left", ARMY_RED, 28, space=8)
        p.paragraph_format.left_indent = Cm(3.0)
        p.paragraph_format.right_indent = Cm(3.0)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        add_text_runs(p, label, base_font=FONT_HEAD, size=Pt(9), bold=True,
                      color=SWAMP_GREEN)
        p.add_run().add_break()
        add_text_runs(p, text, size=Pt(10))
    tail = doc.add_paragraph()
    tail.paragraph_format.space_after = Pt(6)


def add_callout(text_lines):
    p = doc.add_paragraph()
    set_shading(p, MOAWHANGO)
    set_border(p, "left", ARMY_RED, 28, space=8)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    for i, line in enumerate(text_lines):
        if i:
            p.add_run().add_break()
        add_text_runs(p, line, bold=True)
    return p


def add_bullet(text, level=1):
    p = doc.add_paragraph()
    indent = 0.7 * level
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.first_line_indent = Cm(-0.42)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(indent))
    bullet = p.add_run("•\t")
    force_font(bullet, FONT_HEAD)
    force_color(bullet, SWAMP_GREEN)
    bullet.bold = True
    add_text_runs(p, text)
    return p


def add_numbered(num, lead):
    """Numbered point.  Under a Notes heading the type is a little smaller
    and tighter so notes read as supporting material."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.first_line_indent = Cm(-0.7)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(0.7))
    size = None
    if _context["notes"]:
        size = Pt(9.5)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
    else:
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(2)
    n = p.add_run(f"{num}.\t")
    force_font(n, FONT_HEAD)
    if size:
        n.font.size = size
    force_color(n, SWAMP_GREEN if _context["notes"] else DARKEST_HOUR)
    add_text_runs(p, lead, size=size)
    return p


# ------------------------------------------------------- headers & footers --


def marking_paragraph(container, existing=True):
    p = container.paragraphs[0] if existing and container.paragraphs else \
        container.add_paragraph()
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(PROTECTIVE_MARKING)
    force_font(run, FONT_HEAD)
    run.font.size = Pt(10)
    run.bold = True
    force_color(run, DARKEST_HOUR)
    return p


def add_header_footer(section, footer_left, text_width_cm=TEXT_WIDTH_CM):
    marking_paragraph(section.header)
    marking_paragraph(section.footer)
    info = section.footer.add_paragraph()
    info.paragraph_format.space_before = Pt(2)
    info.paragraph_format.space_after = Pt(0)
    info.paragraph_format.tab_stops.add_tab_stop(
        Cm(text_width_cm / 2), WD_TAB_ALIGNMENT.CENTER)
    info.paragraph_format.tab_stops.add_tab_stop(
        Cm(text_width_cm), WD_TAB_ALIGNMENT.RIGHT)

    def footer_run(text):
        r = info.add_run(text)
        force_font(r, FONT_HEAD)
        r.font.size = Pt(8.5)
        force_color(r, DARKEST_HOUR)
        return r

    footer_run(footer_left)
    footer_run("\t")
    footer_run(FOOTER_REFERENCE)
    footer_run("\tPage ")
    for r in add_field(info, "PAGE", "1"):
        force_font(r, FONT_HEAD)
        r.font.size = Pt(8.5)
    footer_run(" of ")
    for r in add_field(info, "NUMPAGES", "1"):
        force_font(r, FONT_HEAD)
        r.font.size = Pt(8.5)


# -------------------------------------------------------------- letterhead --


def add_logo(width_mm=42, space_after=10):
    logo = Image.open(LOGO_FILE).convert("RGBA")
    logo = logo.crop(logo.getchannel("A").getbbox())
    buf = io.BytesIO()
    logo.save(buf, "PNG")
    buf.seek(0)
    doc.add_picture(buf, width=Mm(width_mm))
    logo_para = doc.paragraphs[-1]
    logo_para.paragraph_format.space_before = Pt(0)
    logo_para.paragraph_format.space_after = Pt(space_after)
    return logo_para


def add_locator(current=None, size=Pt(8.5), align_left=True, arrow_color=None):
    """LEAD TEAMS -> LEAD LEADERS -> LEAD SYSTEMS -> COMMAND, with the
    current course in dark green and the rest muted."""
    p = doc.add_paragraph()
    if not align_left:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    for i, step in enumerate(PATHWAY):
        if i:
            a = _run(p, "  →  ", FONT_HEAD, size, False, False,
                     arrow_color or WAIOURU_HILLS)
        here = current is not None and PATHWAY[current] == step
        r = _run(p, step, FONT_HEAD, size, here or current is None, False,
                 SWAMP_GREEN if (here or current is None) else MID_GREY)
        letterspace(r, 15)
    return p


def add_title_block(kicker, title, tag, reference=None, locator=None,
                    with_logo=True, title_style=None, status=None,
                    page_break=False):
    """Layered document opening: small kicker (course code), the title as
    the anchor, a caps tag line closed by the red rule, then originator and
    reference details, then an optional pathway locator."""
    if with_logo:
        add_logo()

    if kicker:
        k = doc.add_paragraph()
        k.paragraph_format.space_after = Pt(0)
        k.paragraph_format.keep_with_next = True
        k.paragraph_format.page_break_before = page_break
        page_break = False
        r = _run(k, kicker.upper(), FONT_HEAD, Pt(10), True, False, SWAMP_GREEN)
        letterspace(r, 40)

    title_p = doc.add_paragraph(style=title_style) if title_style \
        else doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.paragraph_format.keep_with_next = True
    title_p.paragraph_format.page_break_before = page_break
    t = title_p.add_run(title)
    force_font(t, FONT_HEAD)
    t.font.size = Pt(24)
    t.bold = True
    force_color(t, DARKEST_HOUR)

    tag_p = doc.add_paragraph()
    tag_p.paragraph_format.space_after = Pt(6)
    tag_p.paragraph_format.keep_with_next = True
    set_border(tag_p, "bottom", ARMY_RED, 18, space=8)
    if tag:
        r = _run(tag_p, tag.upper(), FONT_HEAD, Pt(9.5), True, False, SWAMP_GREEN)
        letterspace(r, 30)

    org_p = doc.add_paragraph()
    org_p.paragraph_format.space_before = Pt(6)
    org_p.paragraph_format.space_after = Pt(2)
    org_p.paragraph_format.keep_with_next = True
    _run(org_p, ORIGINATOR_LONG, FONT_HEAD, Pt(10), False, False, SWAMP_GREEN)

    if reference:
        meta_p = doc.add_paragraph()
        meta_p.paragraph_format.space_after = Pt(0)
        meta_p.paragraph_format.keep_with_next = True
        parts = [("Reference", reference), ("Date", DATE)]
        if status:
            parts.append(("Status", status))
        for i, (label, value) in enumerate(parts):
            if i:
                _run(meta_p, "   ·   ", FONT_HEAD, Pt(9), False, False, WAIOURU_HILLS)
            _run(meta_p, f"{label}: ", FONT_HEAD, Pt(9), True, False, SWAMP_GREEN)
            _run(meta_p, value, FONT_HEAD, Pt(9), label == "Status", False,
                 ARMY_RED if label == "Status" else DARKEST_HOUR)

    if locator is not None:
        add_locator(locator)
    else:
        gap = doc.add_paragraph()
        gap.paragraph_format.space_after = Pt(10)
    return title_p


def add_letterhead(title, subtitle, reference, with_logo=True,
                   title_style=None, sheet=None, status=None, page_break=False):
    """Course data sheet opening.  `sheet` supplies the code, name and
    pathway position; without it the title is used as-is."""
    if sheet:
        return add_title_block(sheet["code"], sheet["name"], subtitle,
                               reference, locator=PATHWAY_INDEX[sheet["name"]],
                               with_logo=with_logo, title_style=title_style,
                               status=status or STATUS,
                           page_break=page_break)
    return add_title_block(None, title, subtitle, reference,
                           with_logo=with_logo, title_style=title_style,
                           status=status or STATUS,
                           page_break=page_break)


# ------------------------------------------------------------------- body ---


def render_markdown(lines):
    """Typeset the source mini-syntax (see module docstring)."""
    _context["lo"] = False
    _context["notes"] = False
    i = 0
    while i < len(lines):
        line = lines[i]

        if not line.strip():
            i += 1
            continue

        if line.startswith("# "):
            i += 1  # document title already in letterhead
            continue

        if line.startswith("## "):
            heading = add_section_heading(line[3:].strip())
            if line[3:].strip().startswith("Annex"):
                heading.paragraph_format.page_break_before = True
            i += 1
            continue

        if line.startswith("#### "):
            add_sub_heading(line[5:].strip(), 4)
            i += 1
            continue

        if line.startswith("### "):
            add_sub_heading(line[4:].strip(), 3)
            i += 1
            continue

        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(tuple(cells))
                i += 1
            if max(len(r) for r in rows) > 2:
                add_wide_table(rows)
            else:
                add_field_table([(r[0], r[1] if len(r) > 1 else "") for r in rows])
            continue

        if line.startswith(">> "):
            boxes = []
            while i < len(lines) and lines[i].startswith(">> "):
                label, _, text = lines[i][3:].partition("|")
                boxes.append((label.strip(), text.strip()))
                i += 1
            add_chain(boxes)
            continue

        if line.startswith(">! "):
            block = []
            while i < len(lines) and lines[i].startswith(">! "):
                block.append(lines[i][3:].strip())
                i += 1
            add_callout(block)
            continue

        if line.startswith("* "):
            add_bullet(line[2:].strip())
            i += 1
            continue

        if line.startswith("    * "):
            add_bullet(line[6:].strip(), level=2)
            i += 1
            continue

        m = re.match(r"^(\d+)\. (.*)$", line)
        if m:
            add_numbered(m.group(1), m.group(2).strip())
            i += 1
            continue

        add_body(line.strip())
        i += 1


# ----------------------------------------------------------------- finish ---

USED_STYLE_IDS = {"Normal", "Heading1", "Heading2", "DefaultParagraphFont",
                  "TableNormal", "TableGrid", "NoList", "Header", "Footer",
                  "TOC1", "TOC2", "TOCHeading"}


def finish(title):
    styles_el = doc.styles.element
    for st in list(styles_el.findall(qn("w:style"))):
        if st.get(qn("w:styleId")) not in USED_STYLE_IDS:
            styles_el.remove(st)
    for rFonts in styles_el.iter(qn("w:rFonts")):
        for attr in ("ascii", "hAnsi", "cs"):
            rFonts.set(qn(f"w:{attr}"), FONT_BODY)
        for attr in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme",
                     "eastAsia"):
            rFonts.attrib.pop(qn(f"w:{attr}"), None)

    while doc.paragraphs and not doc.paragraphs[-1].text.strip() \
            and not doc.paragraphs[-1]._p.findall(qn("w:r") + "/" + qn("w:drawing")):
        _el = doc.paragraphs[-1]._element
        _el.getparent().remove(_el)

    mark_update_fields(doc)
    props = doc.core_properties
    props.title = title
    props.author = ORIGINATOR


def build_sheet(sheet):
    new_document()
    add_header_footer(doc.sections[0], sheet["footer_left"])
    add_letterhead(sheet["title"], SUBTITLE_LINE, sheet["reference"], sheet=sheet)
    with open(sheet["source"], encoding="utf-8") as fh:
        render_markdown(fh.read().splitlines())
    finish(sheet["title"])
    doc.save(sheet["output"])
    print(f"Saved {sheet['output']}")


if __name__ == "__main__":
    build_sheet(SHEETS[sys.argv[1] if len(sys.argv) > 1 else "lead-teams"])
